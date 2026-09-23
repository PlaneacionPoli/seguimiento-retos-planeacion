"""
Importa BD_Planeacion.xlsx (hojas Retos, Macroactividades_Retos, BD_Cronograma)
a Firestore, siguiendo el modelo de datos de la app (retos + actividades).

Uso:
    python scripts/importar_bd_planeacion.py [ruta_al_excel]

Requiere FIREBASE_SERVICE_ACCOUNT_FILE (o _JSON) en el .env, igual que main.py.
Se puede correr varias veces: cada corrida crea documentos nuevos (no hace upsert),
así que si ya importaste antes, borra las colecciones 'retos' y 'actividades' en la
consola de Firebase antes de reimportar.
"""
import os
import sys
import json
import unicodedata
from pathlib import Path

import openpyxl
import firebase_admin
from firebase_admin import credentials, firestore
from dotenv import load_dotenv

load_dotenv()

EXCEL_PATH = sys.argv[1] if len(sys.argv) > 1 else "BD_Planeacion.xlsx"
UMBRAL_SCORE_MATCH_RETO = 70  # por debajo de esto, se importa sin reto_id para revisión manual


def normalizar(texto):
    if not texto:
        return ""
    texto = str(texto).strip().lower()
    texto = "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")
    return " ".join(texto.split())


def init_firebase():
    service_account_json = os.getenv("FIREBASE_SERVICE_ACCOUNT_JSON")
    service_account_file = os.getenv("FIREBASE_SERVICE_ACCOUNT_FILE", "firebase-service-account.json")

    if service_account_json:
        cred = credentials.Certificate(json.loads(service_account_json))
    elif Path(service_account_file).exists():
        cred = credentials.Certificate(service_account_file)
    else:
        raise SystemExit("No se encontró la service account de Firebase (.env)")

    firebase_admin.initialize_app(cred)
    return firestore.client()


def cargar_usuarios(db):
    """Devuelve {nombre_normalizado: uid} para hacer match de Lider/Responsable por nombre."""
    usuarios = {}
    for doc in db.collection("usuarios").stream():
        data = doc.to_dict()
        nombre = normalizar(data.get("nombre_completo"))
        if nombre:
            usuarios[nombre] = doc.id
    return usuarios


def fecha_iso(valor):
    """Normaliza a 'YYYY-MM-DD'. Algunas hojas (p.ej. Retos) traen la fecha como
    texto con hora ('2026-01-01 00:00:00') en vez de un objeto datetime."""
    if valor is None:
        return None
    try:
        if hasattr(valor, "date"):
            return valor.date().isoformat()
        return str(valor).split(" ")[0].strip() or None
    except Exception:
        return None


def importar_retos(db, wb, usuarios):
    ws = wb["Retos"]
    filas = list(ws.iter_rows(values_only=True))[1:]
    legacy_map = {}  # Reto_ID (excel) -> {"id": doc_id, "proceso":.., "subproceso":..}
    creados = 0

    for fila in filas:
        (reto_id_excel, nombre, proceso, subproceso, area, lider, es_proyecto,
         _num_macro, fecha_inicio, fecha_fin, _cumplimiento) = fila

        if not nombre:
            continue

        lider_id = usuarios.get(normalizar(lider))

        data = {
            "nombre": nombre.strip() if isinstance(nombre, str) else nombre,
            "proceso": proceso,
            "subproceso": subproceso if subproceso and subproceso != "N/A" else None,
            "area": area,
            "lider_id": lider_id,
            "es_proyecto": str(es_proyecto).strip().lower() in ("si", "sí", "yes", "true"),
            "indicador_asociado": None,
            "fecha_inicio": fecha_iso(fecha_inicio),
            "fecha_fin": fecha_iso(fecha_fin),
            "presupuesto_planeado": 0,
            "observaciones": None,
            "created_at": firestore.SERVER_TIMESTAMP,
            "updated_at": firestore.SERVER_TIMESTAMP,
        }
        _, ref = db.collection("retos").add(data)
        legacy_map[reto_id_excel] = {"id": ref.id, "proceso": proceso, "subproceso": data["subproceso"]}
        creados += 1

    print(f"Retos importados: {creados}")
    return legacy_map


def importar_macroactividades(db, wb, legacy_map, usuarios):
    ws = wb["Macroactividades_Retos"]
    filas = list(ws.iter_rows(values_only=True))[1:]
    creadas = 0
    sin_responsable = 0

    for fila in filas:
        (_macro_id_excel, reto_id_excel, _reto_nombre, actividad, entregable,
         fecha_inicio, fecha_fin, avance_real, avance_esperado, _cumplimiento) = fila

        if not actividad:
            continue

        reto_info = legacy_map.get(reto_id_excel)
        if not reto_info:
            continue

        data = {
            "reto_id": reto_info["id"],
            "parent_id": None,
            "es_macroactividad": True,
            "proceso": reto_info["proceso"],
            "subproceso": reto_info["subproceso"],
            "nombre": actividad.strip() if isinstance(actividad, str) else actividad,
            "descripcion": None,
            "entregable": entregable,
            "responsable_id": None,
            "fecha_inicio": fecha_iso(fecha_inicio),
            "fecha_fin": fecha_iso(fecha_fin),
            "periodicidad": None,
            "prioridad": "media",
            "estado": "Completado" if (avance_real or 0) >= 100 else ("En curso" if (avance_real or 0) > 0 else "No iniciado"),
            "avance_actual": avance_real or 0,
            "avance_esperado": avance_esperado or 100,
            "presupuesto_planeado": 0,
            "presupuesto_ejecutado": 0,
            "evidencia_url": None,
            "observaciones": None,
            "comentario_actual": None,
            "orden": 0,
            "tags": [],
            "created_at": firestore.SERVER_TIMESTAMP,
            "updated_at": firestore.SERVER_TIMESTAMP,
        }
        db.collection("actividades").add(data)
        creadas += 1
        sin_responsable += 1  # las macroactividades del Excel no traen responsable individual

    print(f"Macroactividades importadas: {creadas} (todas sin responsable asignado, se asigna en Administración si aplica)")


def importar_cronograma(db, wb, legacy_map, usuarios):
    ws = wb["BD_Cronograma"]
    filas = list(ws.iter_rows(values_only=True))[1:]
    creadas = 0
    sin_responsable = 0
    sin_reto_por_score_bajo = 0

    for fila in filas:
        (_id, proceso, subproceso, actividad, fecha_inicio, fecha_fin, _trimestre, _anio,
         responsable, reto_id_excel, _reto_nombre, score_match, observaciones,
         _fila_origen, _celdas_origen) = fila

        if not actividad:
            continue

        responsable_id = usuarios.get(normalizar(responsable))
        if not responsable_id:
            sin_responsable += 1

        reto_id = None
        if reto_id_excel and legacy_map.get(reto_id_excel):
            if score_match is None or score_match >= UMBRAL_SCORE_MATCH_RETO:
                reto_id = legacy_map[reto_id_excel]["id"]
            else:
                sin_reto_por_score_bajo += 1

        data = {
            "reto_id": reto_id,
            "parent_id": None,
            "es_macroactividad": False,
            "proceso": proceso,
            "subproceso": subproceso if subproceso and subproceso != "N/A" else None,
            "nombre": actividad.strip() if isinstance(actividad, str) else actividad,
            "descripcion": None,
            "entregable": None,
            "responsable_id": responsable_id,
            "fecha_inicio": fecha_iso(fecha_inicio),
            "fecha_fin": fecha_iso(fecha_fin),
            "periodicidad": None,
            "prioridad": "media",
            "estado": "No iniciado",
            "avance_actual": 0,
            "avance_esperado": 100,
            "presupuesto_planeado": 0,
            "presupuesto_ejecutado": 0,
            "evidencia_url": None,
            "observaciones": observaciones,
            "comentario_actual": None,
            "orden": 0,
            "tags": [],
            "created_at": firestore.SERVER_TIMESTAMP,
            "updated_at": firestore.SERVER_TIMESTAMP,
        }
        db.collection("actividades").add(data)
        creadas += 1

    print(f"Actividades de cronograma importadas: {creadas}")
    print(f"  - sin responsable emparejado por nombre: {sin_responsable} (asignar manualmente en Administración)")
    print(f"  - con Reto_ID en el Excel pero descartado por score de match bajo: {sin_reto_por_score_bajo}")


def main():
    if not Path(EXCEL_PATH).exists():
        raise SystemExit(f"No se encontró el archivo: {EXCEL_PATH}")

    db = init_firebase()
    usuarios = cargar_usuarios(db)
    print(f"Usuarios existentes para hacer match por nombre: {len(usuarios)}")

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    legacy_map = importar_retos(db, wb, usuarios)
    importar_macroactividades(db, wb, legacy_map, usuarios)
    importar_cronograma(db, wb, legacy_map, usuarios)

    print("\nImportación completa.")


if __name__ == "__main__":
    main()
